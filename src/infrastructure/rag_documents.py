"""PDF/DOCX/HTML/JSON 文档解析与结构化字符切分。"""
import hashlib
import json
import os
import re
from collections.abc import Iterable

import fitz
from bs4 import BeautifulSoup
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from src.core.settings import BASE_DIR, resolve_project_path, settings
from src.infrastructure.rag_sources import MANIFEST_PATH, load_manifest

KNOWLEDGE_DIR = os.path.join(BASE_DIR, "app", "agent", "rag", "knowledge")
_SUPPORTED_LOCAL = {".pdf", ".docx", ".html", ".htm", ".json"}
_PDF_NOISE = re.compile(r"^(?:Page\s+\d+|DICOM PS3\.|Copyright ©|[-–— ]*Standard[-–— ]*)", re.I)
_SECTION = re.compile(r"^(?:[A-Z]{1,3}(?:\.\d+)*|\d+(?:\.\d+)*)\s+\S+")
_CHUNK_SIZE = 800
_CHUNK_OVERLAP = 120


def _clean_text(text: str) -> str:
    paragraphs = []
    for raw_paragraph in re.split(r"\n\s*\n", text or ""):
        lines = [re.sub(r"\s+", " ", line).strip() for line in raw_paragraph.splitlines()]
        paragraph = "\n".join(
            line for line in lines if line and not _PDF_NOISE.match(line)
        ).strip()
        if paragraph:
            paragraphs.append(paragraph)
    return "\n\n".join(paragraphs)


def _sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fp:
        for block in iter(lambda: fp.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _pdf_units(path: str) -> list[dict]:
    units, paragraphs = [], []
    section = ""
    page_start = None
    page_end = None

    def append_paragraph(lines: list[str], page_no: int):
        nonlocal page_start, page_end
        paragraph = _clean_text("\n".join(lines))
        if not paragraph:
            return
        paragraphs.append(paragraph)
        page_start = page_start or page_no
        page_end = page_no

    def flush():
        nonlocal page_start, page_end
        if paragraphs:
            units.append({
                "text": "\n\n".join(paragraphs),
                "metadata": {
                    "page_start": page_start,
                    "page_end": page_end,
                    "section": section,
                },
            })
            paragraphs.clear()
        page_start = None
        page_end = None

    with fitz.open(path) as document:
        for page_no, page in enumerate(document, start=1):
            blocks = page.get_text("blocks", sort=True)
            if not blocks:
                continue
            for block in blocks:
                lines = [line for line in _clean_text(block[4]).splitlines() if line]
                body_lines: list[str] = []
                for line in lines:
                    if _SECTION.match(line):
                        append_paragraph(body_lines, page_no)
                        body_lines = []
                        flush()
                        section = line
                        page_start = page_no
                        page_end = page_no
                    else:
                        body_lines.append(line)
                append_paragraph(body_lines, page_no)

            # 没有识别出章节时按页形成 unit；有章节时延续到后续页面的新章节。
            if not section:
                flush()
    flush()
    if not units:
        raise ValueError("PDF 未提取到文本，可能是扫描文档（首版不支持 OCR）")
    return units


def _docx_units(path: str) -> list[dict]:
    document = Document(path)
    units: list[dict] = []
    headings: list[str] = []
    buffer: list[str] = []

    def flush():
        if buffer:
            units.append({"text": "\n\n".join(buffer), "metadata": {
                "section": " > ".join(headings),
            }})
            buffer.clear()

    for paragraph in document.paragraphs:
        text = _clean_text(paragraph.text)
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            flush()
            level_match = re.search(r"(\d+)$", paragraph.style.name)
            level = int(level_match.group(1)) if level_match else 1
            headings[level - 1:] = [text]
        else:
            buffer.append(text)
    flush()
    for table in document.tables:
        rows = [" | ".join(_clean_text(cell.text) for cell in row.cells) for row in table.rows]
        text = "\n".join(row for row in rows if row.strip(" |"))
        if text:
            units.append({"text": text, "metadata": {"section": " > ".join(headings)}})
    return units


def _html_units(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as fp:
        soup = BeautifulSoup(fp.read(), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    units: list[dict] = []
    headings: list[str] = []
    buffer: list[str] = []

    def flush():
        if buffer:
            units.append({
                "text": "\n\n".join(buffer),
                "metadata": {"section": " > ".join(headings)},
            })
            buffer.clear()

    root = soup.find("main") or soup.find(role="main") or soup.body or soup
    for element in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "pre", "table"]):
        text = _clean_text(element.get_text(" | " if element.name == "table" else " ", strip=True))
        if not text:
            continue
        if element.name.startswith("h"):
            flush()
            level = int(element.name[1])
            headings[level - 1:] = [text]
        else:
            buffer.append(text)
    flush()
    return units


def _json_units(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as fp:
        raw = json.load(fp)
    rows = raw if isinstance(raw, list) else [raw]
    return [{"id": row.get("id"), "text": row["content"],
             "metadata": dict(row.get("metadata") or {})} for row in rows]


def _load_units(path: str, source_type: str) -> list[dict]:
    source_type = source_type.lower()
    if source_type == "pdf":
        return _pdf_units(path)
    if source_type == "docx":
        return _docx_units(path)
    if source_type in ("html", "htm"):
        return _html_units(path)
    if source_type == "json":
        return _json_units(path)
    raise ValueError("不支持的知识文档类型: %s" % source_type)


def _split_units(units: Iterable[dict], base_metadata: dict) -> list[dict]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=_CHUNK_SIZE,
        chunk_overlap=_CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", "。", ". ", "；", "; ", " ", ""],
    )
    chunks = []
    for unit_index, unit in enumerate(units):
        content = _clean_text(unit["text"])
        if not content:
            continue
        parts = splitter.split_text(content)
        for index, part in enumerate(parts):
            content_hash = hashlib.sha256(part.encode("utf-8")).hexdigest()
            metadata = {
                "doc_id": "",
                "document_hash": "",
                "category": "uncategorized",
                "source_url": "",
                "version": "",
                "title": "",
                "section": "",
                "page_start": None,
                "page_end": None,
                **base_metadata,
                **(unit.get("metadata") or {}),
                "unit_index": unit_index,
                "chunk_index": index,
                "content_hash": content_hash,
            }
            seed = "|".join(str(x) for x in (
                base_metadata.get("doc_id", base_metadata.get("source_id", "")),
                metadata.get("section", ""),
                metadata.get("page_start", ""), unit.get("id", ""), unit_index,
                index, content_hash,
            ))
            chunk_id = unit.get("id") if len(parts) == 1 and unit.get("id") else (
                "chunk-" + hashlib.sha256(seed.encode("utf-8")).hexdigest()[:24]
            )
            chunks.append({"id": chunk_id, "content": part, "metadata": metadata})
    return chunks


def document_manifest(manifest_path: str = MANIFEST_PATH) -> list[dict]:
    """返回当前知识文档快照；只读取文件字节计算 hash，不解析正文。"""
    source_dir = resolve_project_path(settings.rag.source_dir)
    documents: list[dict] = []

    for name in sorted(os.listdir(KNOWLEDGE_DIR)):
        if not name.endswith(".json"):
            continue
        path = os.path.join(KNOWLEDGE_DIR, name)
        digest = _sha256(path)
        documents.append({
            "doc_id": "builtin-" + os.path.splitext(name)[0],
            "path": path,
            "type": "json",
            "hash": digest,
            "sha256": digest,
            "metadata": {
                "source_id": "builtin-" + os.path.splitext(name)[0],
                "doc_id": "builtin-" + os.path.splitext(name)[0],
                "title": name,
                "version": "builtin",
                "source_type": "json",
                "source_url": "",
            },
        })

    for source in load_manifest(manifest_path):
        path = os.path.join(source_dir, os.path.basename(source["filename"]))
        if not os.path.isfile(path):
            raise FileNotFoundError("RAG source not downloaded: %s" % path)
        digest = _sha256(path)
        documents.append({
            "doc_id": source["id"],
            "path": path,
            "type": source["type"],
            "hash": digest,
            "sha256": digest,
            "metadata": {
                "source_id": source["id"],
                "doc_id": source["id"],
                "title": source["title"],
                "version": str(source["version"]),
                "source_type": source["type"],
                "source_url": source["url"],
                "category": source["category"],
            },
        })

    local_dir = os.path.join(source_dir, "local")
    if os.path.isdir(local_dir):
        for root, _, names in os.walk(local_dir):
            for name in sorted(names):
                ext = os.path.splitext(name)[1].lower()
                if ext not in _SUPPORTED_LOCAL:
                    continue
                path = os.path.join(root, name)
                relative = os.path.relpath(path, local_dir).replace(os.sep, "/")
                doc_id = "local-" + hashlib.sha256(relative.encode("utf-8")).hexdigest()[:16]
                digest = _sha256(path)
                documents.append({
                    "doc_id": doc_id,
                    "path": path,
                    "type": ext.lstrip("."),
                    "hash": digest,
                    "sha256": digest,
                    "metadata": {
                        "source_id": doc_id,
                        "doc_id": doc_id,
                        "title": name,
                        "version": "local",
                        "source_type": ext.lstrip("."),
                        "source_url": "",
                        "category": "local_doc",
                    },
                })
    return documents


def build_chunks(
    manifest_path: str = MANIFEST_PATH,
    document_ids: set[str] | None = None,
) -> list[dict]:
    """解析指定文档；document_ids=None 时解析全部文档。"""
    chunks = []
    for document in document_manifest(manifest_path):
        if document_ids is not None and document["doc_id"] not in document_ids:
            continue
        metadata = {**document["metadata"], "document_hash": document["hash"]}
        units = _load_units(document["path"], document["type"])
        chunks.extend(_split_units(units, metadata))
    return chunks

